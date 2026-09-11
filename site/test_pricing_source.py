"""Tests for the cost workbook parser.

Run: uv run --with pytest --with openpyxl --with weasyprint==70.0 pytest site/test_pricing_source.py -q
"""

import os
import sys
from decimal import Decimal, ROUND_HALF_UP, ROUND_UP
from pathlib import Path

import openpyxl
import pytest

SITE_DIR = Path(__file__).resolve().parent
REPO_ROOT = SITE_DIR.parent
RECIPES_CAFE = Path(
    os.environ.get("RECIPES_CAFE", REPO_ROOT.parent / "recipes" / "cafe.md")
)
CAFE_COSTS_XLSX = RECIPES_CAFE.parent / "cafe_costs.xlsx"

sys.path.insert(0, str(SITE_DIR))

import pricing_source  # noqa: E402
import generate  # noqa: E402

sys.path.insert(0, str(REPO_ROOT / "menu"))

import menu_source  # noqa: E402

# Spot values are pinned to one millionth: the evaluator must agree with the
# hand-derived arithmetic to full displayed-cents precision and far beyond.
COST_ABS_TOLERANCE = 1e-6

# The charm pricing contract lives at fixed workbook addresses: these Read
# Me cells, the Drinks multiplier column, and the minus-a-cent charm
# adjustment. A workbook restructure that moves any of them must update the
# constant alongside the mirrored formulas below.
MARGIN_MULTIPLIER_CELL = "B12"
OPERATING_COST_CELL = "B13"
DRINK_PRICE_MULTIPLIER_COLUMN = "E"
CHARM_CENT_ADJUSTMENT = 0.01


def read_rows_by_name() -> dict[str, pricing_source.DrinkCost]:
    return {
        row.name: row for row in pricing_source.read_drink_costs(CAFE_COSTS_XLSX)
    }


def open_costs_workbook() -> openpyxl.Workbook:
    """Open the live cost workbook for one read of owner-curated values."""

    return openpyxl.load_workbook(CAFE_COSTS_XLSX)


def read_readme_value(workbook: openpyxl.Workbook, coordinate: str) -> float:
    """Return one numeric Read Me cell from an opened cost workbook.

    The pricing parameter is read at test time because the workbook is
    owner-curated live data: pinning its value in the suite would break on
    every legitimate workbook edit.
    """

    return float(workbook["Read Me"][coordinate].value)


def read_drink_cell(
    workbook: openpyxl.Workbook, drink_name: str, column: str
) -> float:
    """Return one numeric Drinks-sheet cell for a named drink from an
    opened cost workbook.

    Like the Read Me parameters, the per-drink Price Multiplier is read
    at test time because the workbook is owner-curated live data.
    """

    sheet = workbook[pricing_source.DRINKS_SHEET]
    for row in range(pricing_source.FIRST_DATA_ROW, sheet.max_row + 1):
        if sheet[f"{pricing_source.DRINK_NAME_COLUMN}{row}"].value == drink_name:
            return float(sheet[f"{column}{row}"].value)
    raise AssertionError(f"no {drink_name!r} row on the Drinks sheet")


def round_half_away(value: float, digits: int = 0) -> float:
    """Round half away from zero to ``digits`` places, Excel ROUND.

    Written here rather than imported from the module under test so the
    pin's arithmetic stays independent; Python's round() is half-even and
    cannot serve.
    """

    quantum = Decimal(1).scaleb(-digits)
    return float(Decimal(str(value)).quantize(quantum, rounding=ROUND_HALF_UP))


def roundup_away(value: float, digits: int = 0) -> float:
    """Round any remainder away from zero to ``digits`` places, Excel
    ROUNDUP."""

    quantum = Decimal(1).scaleb(-digits)
    return float(Decimal(str(value)).quantize(quantum, rounding=ROUND_UP))


def expected_charm_price(drink_name: str, ingredient_cost: float) -> tuple[float, float]:
    """Return the charm and floor terms of the Menu Price contract for a
    drink: the rounded markup minus a cent, and the lowest .99 at or above
    Total Cost.

    The margin and operating costs and the drink's per-drink multiplier
    are read live: the workbook is owner-curated, so pinning them would
    break on every legitimate workbook edit.
    """

    workbook = open_costs_workbook()
    margin_multiplier = read_readme_value(workbook, MARGIN_MULTIPLIER_CELL)
    operating_cost = read_readme_value(workbook, OPERATING_COST_CELL)
    price_multiplier = read_drink_cell(
        workbook, drink_name, DRINK_PRICE_MULTIPLIER_COLUMN
    )
    charm = round_half_away(
        ingredient_cost * margin_multiplier * price_multiplier + operating_cost, 0
    ) - CHARM_CENT_ADJUSTMENT
    floor = roundup_away(ingredient_cost + operating_cost, 0) - CHARM_CENT_ADJUSTMENT
    return charm, floor


def write_synthetic_workbook(
    path: Path,
    cost_formula: str,
    multiplier: float = 2.0,
    operating_cost: float = 0.0,
) -> Path:
    """Write a minimal workbook mirroring the real sheet shape.

    The Drinks C2 cell carries ``cost_formula`` as the ingredient-cost
    stand-in; D2 adds the Read Me B13 operating cost flat onto it, and
    F2 holds the workbook's charm formula over C2, D2, the Read Me B12
    multiplier, and the per-drink E2 multiplier. Bases D5 and
    Ingredients E2 give every drink formula a two-link resolution chain
    through other formula cells.
    """

    workbook = openpyxl.Workbook()
    readme = workbook.active
    readme.title = "Read Me"
    readme[MARGIN_MULTIPLIER_CELL] = multiplier
    readme[OPERATING_COST_CELL] = operating_cost
    ingredients = workbook.create_sheet("Ingredients")
    ingredients["A2"] = "Turbinado Sugar"
    ingredients["C2"] = 8.49
    ingredients["D2"] = 1814
    ingredients["E2"] = "=C2/D2"
    bases = workbook.create_sheet("Bases")
    bases["A5"] = "Turbinado Syrup (25g dose)"
    bases["D5"] = "=Ingredients!E2/10"
    drinks = workbook.create_sheet("Drinks")
    drinks["A2"] = "Test Drink"
    drinks["C2"] = cost_formula
    drinks["D2"] = "=C2+'Read Me'!$B$13"
    drinks["E2"] = 1
    drinks["F2"] = (
        "=MAX(ROUND(C2*'Read Me'!$B$12*E2+'Read Me'!$B$13,0)-0.01,"
        "ROUNDUP(D2,0)-0.01)"
    )
    workbook.save(path)
    return path


class TestReadDrinkCosts:
    def test_read_drink_costs_returns_all_rows_from_workbook(self):
        rows = pricing_source.read_drink_costs(CAFE_COSTS_XLSX)

        # The expected count derives from the Drinks sheet at test time:
        # the owner adds rows as live data (29 grew to 32 mid-build), so
        # pinning a count would break on every legitimate edit.
        workbook = openpyxl.load_workbook(CAFE_COSTS_XLSX)
        sheet = workbook[pricing_source.DRINKS_SHEET]
        sheet_row_count = sum(
            1
            for row in range(pricing_source.FIRST_DATA_ROW, sheet.max_row + 1)
            if sheet[f"{pricing_source.DRINK_NAME_COLUMN}{row}"].value is not None
        )
        assert len(rows) == sheet_row_count
        assert len(rows) > 0

    def test_read_drink_costs_preserves_workbook_row_order(self):
        rows = pricing_source.read_drink_costs(CAFE_COSTS_XLSX)

        # The expected order derives from the Drinks sheet at test time:
        # the owner reorders and renames rows as live data, so pinning
        # names would break on every legitimate edit.
        workbook = openpyxl.load_workbook(CAFE_COSTS_XLSX)
        sheet = workbook[pricing_source.DRINKS_SHEET]
        sheet_names = [
            sheet[f"{pricing_source.DRINK_NAME_COLUMN}{row}"].value
            for row in range(
                pricing_source.FIRST_DATA_ROW, sheet.max_row + 1
            )
            if sheet[f"{pricing_source.DRINK_NAME_COLUMN}{row}"].value
            is not None
        ]
        assert [row.name for row in rows] == sheet_names

    def test_black_coffee_cost_matches_hand_derived_total_cost(self):
        # The workbook's Total Cost relationship is owner-curated: the
        # D column adds the Read Me operating cost flat onto the
        # ingredient cost. The operating cost value is read live; a
        # future structural change to the D-column formulas must update
        # this pin's shape.
        operating_cost = read_readme_value(open_costs_workbook(), OPERATING_COST_CELL)

        black_coffee = read_rows_by_name()["Black Coffee"]

        assert black_coffee.cost == pytest.approx(
            20 * 18.99 / 1134 + operating_cost, abs=COST_ABS_TOLERANCE
        )

    def test_black_coffee_suggested_matches_workbook_charm_pricing(self):
        # The workbook's Menu Price relationship is owner-curated: the
        # F column rounds the marked-up ingredient cost to the nearest
        # whole dollar minus a cent but never below the lowest .99 at or
        # above Total Cost (the Read Me C12 charm contract); a future
        # structural change to the F-column formulas must update this
        # pin's shape.
        black_coffee = read_rows_by_name()["Black Coffee"]

        charm, floor = expected_charm_price("Black Coffee", 20 * 18.99 / 1134)

        assert black_coffee.suggested == pytest.approx(
            max(charm, floor), abs=COST_ABS_TOLERANCE
        )

    def test_house_latte_cost_matches_hand_derived_build(self):
        # 1 batch concentrate + 25g syrup dose + 150g whole milk, straight
        # from the workbook's Bases D2, Bases D5, and Ingredients E3 chains,
        # plus the flat operating cost the Total Cost column adds.
        operating_cost = read_readme_value(open_costs_workbook(), OPERATING_COST_CELL)

        house_latte = read_rows_by_name()["House Latte"]

        ingredient_cost = (
            20 * 18.99 / 1134
            + (100 * 8.49 / 1814) / 250 * 25
            + 150 * 5.99 / 7570
        )
        assert house_latte.cost == pytest.approx(
            ingredient_cost + operating_cost, abs=COST_ABS_TOLERANCE
        )

    def test_house_latte_suggested_matches_workbook_charm_pricing(self):
        house_latte = read_rows_by_name()["House Latte"]

        ingredient_cost = (
            20 * 18.99 / 1134
            + (100 * 8.49 / 1814) / 250 * 25
            + 150 * 5.99 / 7570
        )
        charm, floor = expected_charm_price("House Latte", ingredient_cost)

        assert house_latte.suggested == pytest.approx(
            max(charm, floor), abs=COST_ABS_TOLERANCE
        )

    def test_hot_tea_cost_matches_hand_derived_tea_weight(self):
        operating_cost = read_readme_value(open_costs_workbook(), OPERATING_COST_CELL)

        hot_tea = read_rows_by_name()["Hot Tea"]

        assert hot_tea.cost == pytest.approx(
            4 * 14.99 / 454 + operating_cost, abs=COST_ABS_TOLERANCE
        )

    def test_hot_tea_suggested_floors_at_the_charm_price_over_cost(self):
        # Hot Tea's marked-up ingredient cost rounds to $4.99, below its
        # Total Cost, so this row exercises the charm contract's floor
        # rule: the Menu Price must be the lowest .99 at or above Total
        # Cost, not the rounded markup.
        hot_tea = read_rows_by_name()["Hot Tea"]

        charm, floor = expected_charm_price("Hot Tea", 4 * 14.99 / 454)

        assert charm < floor, (
            "precondition: this row's rounded markup must sit below the "
            "floor for the floor rule to be exercised"
        )
        assert hot_tea.suggested == pytest.approx(floor, abs=COST_ABS_TOLERANCE)

    def test_every_suggested_price_formats_to_the_charm_99_ending(self):
        # The workbook Read Me documents the Menu Price contract as
        # rounded to the nearest .99; this pins the contract across every
        # row through the same formatter the pages display with.
        off_contract = [
            (row.name, generate.format_price(row.suggested))
            for row in pricing_source.read_drink_costs(CAFE_COSTS_XLSX)
            if not generate.format_price(row.suggested).endswith(".99")
        ]
        assert off_contract == []


class TestFormulaEvaluation:
    def test_b12_multiplier_edit_reprices_suggested_at_next_read(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=Ingredients!E2*100"
        )
        base_multiplier = 2.0
        edited_multiplier = 3.5
        ingredient_cost = (8.49 / 1814) * 100
        floor = roundup_away(ingredient_cost, 0) - CHARM_CENT_ADJUSTMENT

        before = pricing_source.read_drink_costs(path)[0]
        before_charm = (
            round_half_away(ingredient_cost * base_multiplier, 0)
            - CHARM_CENT_ADJUSTMENT
        )
        assert before.cost == pytest.approx(
            ingredient_cost, abs=COST_ABS_TOLERANCE
        )
        assert before.suggested == pytest.approx(
            max(before_charm, floor), abs=COST_ABS_TOLERANCE
        )

        workbook = openpyxl.load_workbook(path)
        workbook["Read Me"][MARGIN_MULTIPLIER_CELL] = edited_multiplier
        workbook.save(path)

        after = pricing_source.read_drink_costs(path)[0]
        after_charm = (
            round_half_away(ingredient_cost * edited_multiplier, 0)
            - CHARM_CENT_ADJUSTMENT
        )
        assert after.suggested == pytest.approx(
            max(after_charm, floor), abs=COST_ABS_TOLERANCE
        )

    def test_formula_chain_resolves_through_nested_cells(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=Bases!D5*2"
        )

        (row,) = pricing_source.read_drink_costs(path)

        assert row.cost == pytest.approx(
            (8.49 / 1814 / 10) * 2, abs=COST_ABS_TOLERANCE
        )

    def test_function_call_formula_fails_naming_cell_and_formula(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=SUM(A1:A2)"
        )

        with pytest.raises(pricing_source.FormulaError) as excinfo:
            pricing_source.read_drink_costs(path)

        message = str(excinfo.value)
        assert "Drinks!C2" in message
        assert "=SUM(A1:A2)" in message

    def test_bare_word_formula_fails_naming_cell_and_formula(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=Bogus+1"
        )

        with pytest.raises(pricing_source.FormulaError) as excinfo:
            pricing_source.read_drink_costs(path)

        message = str(excinfo.value)
        assert "Drinks!C2" in message
        assert "=Bogus+1" in message

    def test_self_reference_fails_naming_the_cyclic_cell(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=C2+1"
        )

        with pytest.raises(pricing_source.FormulaError) as excinfo:
            pricing_source.read_drink_costs(path)

        message = str(excinfo.value)
        assert "Drinks!C2" in message
        assert "cycle" in message.lower()

    def test_reference_to_empty_cell_fails_naming_the_reference(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=Z99*2"
        )

        with pytest.raises(pricing_source.FormulaError) as excinfo:
            pricing_source.read_drink_costs(path)

        assert "Z99" in str(excinfo.value)


class TestExcelFunctionFormulas:
    """Focused parser tests for the three supported Excel functions.

    The workbook's Menu Price column calls MAX, ROUND, and ROUNDUP with
    Excel rounding semantics, which break ties away from zero where
    Python's round() is half-even.
    """

    def test_max_returns_the_largest_argument(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=MAX(1.5,Bases!D5,2.25)"
        )

        (row,) = pricing_source.read_drink_costs(path)

        assert row.cost == pytest.approx(2.25, abs=COST_ABS_TOLERANCE)

    def test_round_breaks_a_tie_away_from_zero_not_half_even(self, tmp_path):
        # Excel ROUND(4.5,0) is 5 while Python's half-even round(4.5) is
        # 4, so this value pins the Excel semantics specifically.
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=ROUND(4.5,0)"
        )

        (row,) = pricing_source.read_drink_costs(path)

        assert row.cost == pytest.approx(5.0, abs=COST_ABS_TOLERANCE)

    def test_round_digits_beyond_zero_round_half_away(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=ROUND(0.345,2)"
        )

        (row,) = pricing_source.read_drink_costs(path)

        assert row.cost == pytest.approx(0.35, abs=COST_ABS_TOLERANCE)

    def test_round_rounds_a_negative_tie_away_from_zero(self, tmp_path):
        # The grammar has no unary minus, so the negative value arrives
        # as a subtraction; Excel ROUND(-5.5,0) is -6, away from zero.
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=ROUND(0-5.5,0)"
        )

        (row,) = pricing_source.read_drink_costs(path)

        assert row.cost == pytest.approx(-6.0, abs=COST_ABS_TOLERANCE)

    def test_roundup_pushes_any_remainder_away_from_zero(self, tmp_path):
        # Excel ROUNDUP(5.1,0) is 6; a nearest-value rounding would say 5.
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=ROUNDUP(5.1,0)"
        )

        (row,) = pricing_source.read_drink_costs(path)

        assert row.cost == pytest.approx(6.0, abs=COST_ABS_TOLERANCE)

    def test_roundup_keeps_an_already_round_value(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=ROUNDUP(5,0)"
        )

        (row,) = pricing_source.read_drink_costs(path)

        assert row.cost == pytest.approx(5.0, abs=COST_ABS_TOLERANCE)

    def test_roundup_rounds_a_negative_remainder_away_from_zero(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=ROUNDUP(0-5.1,0)"
        )

        (row,) = pricing_source.read_drink_costs(path)

        assert row.cost == pytest.approx(-6.0, abs=COST_ABS_TOLERANCE)

    @pytest.mark.parametrize(
        "formula", ["=MAX(4.5)", "=ROUND(4.5)", "=ROUNDUP(4.5,0,1)"]
    )
    def test_wrong_function_arity_fails_naming_cell_and_formula(
        self, tmp_path, formula
    ):
        path = write_synthetic_workbook(tmp_path / "costs.xlsx", cost_formula=formula)

        with pytest.raises(pricing_source.FormulaError) as excinfo:
            pricing_source.read_drink_costs(path)

        message = str(excinfo.value)
        assert "Drinks!C2" in message
        assert formula in message

    def test_unknown_function_name_fails_naming_cell_and_formula(self, tmp_path):
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=SUM(1,2)"
        )

        with pytest.raises(pricing_source.FormulaError) as excinfo:
            pricing_source.read_drink_costs(path)

        message = str(excinfo.value)
        assert "Drinks!C2" in message
        assert "=SUM(1,2)" in message
        assert "unknown function" in message

    def test_round_overflow_fails_as_a_formula_error(self, tmp_path):
        # A magnitude beyond the evaluator's decimal precision must fail
        # loudly like any other unevaluable formula, not escape as a raw
        # decimal exception.
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx",
            cost_formula="=ROUND(10000000000000000000000000000,0)",
        )

        with pytest.raises(pricing_source.FormulaError) as excinfo:
            pricing_source.read_drink_costs(path)

        message = str(excinfo.value)
        assert "Drinks!C2" in message
        assert "precision" in message

    def test_round_digits_overflowing_to_infinity_fail_as_a_formula_error(
        self, tmp_path
    ):
        # Digits that overflow the evaluator's float arithmetic to
        # infinity (1e160 * 1e160) reach the rounding quantum's int()
        # conversion, which must fail as a FormulaError naming the cell
        # like every other unevaluable formula, not escape as a raw
        # OverflowError.
        overflow_digits = "1" + "0" * 160
        formula = f"=ROUND(4.5,{overflow_digits}*{overflow_digits})"
        path = write_synthetic_workbook(tmp_path / "costs.xlsx", cost_formula=formula)

        with pytest.raises(pricing_source.FormulaError) as excinfo:
            pricing_source.read_drink_costs(path)

        message = str(excinfo.value)
        assert "Drinks!C2" in message
        assert "precision" in message


class TestContractSurface:
    def test_price_aliases_match_contract(self):
        assert pricing_source.PRICE_ALIASES == {
            "Pour Over": "Pour Over Coffee",
            "Vietnamese Coffee": "Vietnamese Iced Coffee",
        }

    def test_bac_xiu_item_matches_contract(self):
        item = pricing_source.BAC_XIU

        assert isinstance(item, menu_source.Item)
        assert item.name_en == "Bạc Xỉu"
        assert item.name_vi == "Cà Phê Bạc Xỉu"
        assert item.description == (
            "The classic milkier white coffee: the House Latte built on "
            "condensed milk instead of syrup."
        )
        assert item.temperatures == ["hot", "iced"]

    def test_bac_xiu_anchor_matches_contract(self):
        assert pricing_source.BAC_XIU_ANCHOR == "House Latte"

    def test_default_costs_derives_from_sibling_recipes_checkout(self):
        assert pricing_source.DEFAULT_COSTS == (
            REPO_ROOT.parent / "recipes" / "cafe_costs.xlsx"
        )


class TestLazyOpenpyxlImport:
    def test_missing_openpyxl_raises_runtime_error_naming_uv_incantation(
        self, tmp_path, monkeypatch
    ):
        monkeypatch.setitem(sys.modules, "openpyxl", None)
        path = write_synthetic_workbook(
            tmp_path / "costs.xlsx", cost_formula="=Ingredients!E2"
        )

        with pytest.raises(RuntimeError) as excinfo:
            pricing_source.read_drink_costs(path)

        assert "uv run --with openpyxl python site/generate.py" in str(
            excinfo.value
        )
